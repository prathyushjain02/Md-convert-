"""End-to-end tests against the Flask app with real MarkItDown conversions."""

from __future__ import annotations

import io
import json
import zipfile

import pytest

from app import create_app
from app.converter import markdown_filename_for, safe_stem


@pytest.fixture()
def client():
    app = create_app()
    app.config.update(TESTING=True)
    with app.test_client() as test_client:
        yield test_client


def upload(name: str, data: bytes) -> dict:
    return {"files": (io.BytesIO(data), name)}


def test_index_renders(client):
    response = client.get("/")
    assert response.status_code == 200
    assert b"Turn any document into Markdown" in response.data


def test_healthz(client):
    response = client.get("/healthz")
    assert response.status_code == 200
    assert response.get_json() == {"status": "ok"}


def test_limits_endpoint(client):
    payload = client.get("/api/limits").get_json()
    assert payload["max_files"] >= 1
    assert ".pdf" in payload["allowed_extensions"]


def test_convert_html_to_markdown(client):
    html = b"<html><body><h1>Title</h1><p>Hello <b>world</b>.</p></body></html>"
    response = client.post(
        "/api/convert", data=upload("page.html", html), content_type="multipart/form-data"
    )
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["summary"] == {"total": 1, "succeeded": 1, "failed": 0}

    result = payload["results"][0]
    assert result["ok"] is True
    assert result["markdown_filename"] == "page.md"
    assert "# Title" in result["markdown"]
    assert "**world**" in result["markdown"]


def test_convert_csv_to_markdown_table(client):
    csv = b"name,role\nAda,engineer\nGrace,admiral\n"
    payload = client.post(
        "/api/convert", data=upload("team.csv", csv), content_type="multipart/form-data"
    ).get_json()

    markdown = payload["results"][0]["markdown"]
    assert "| name | role |" in markdown
    assert "Grace" in markdown


def test_convert_docx_to_markdown(client):
    docx = pytest.importorskip("docx", reason="python-docx not installed")  # noqa: F841
    from docx import Document

    document = Document()
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("Revenue grew.")
    buffer = io.BytesIO()
    document.save(buffer)

    payload = client.post(
        "/api/convert",
        data=upload("Quarterly Report.docx", buffer.getvalue()),
        content_type="multipart/form-data",
    ).get_json()

    result = payload["results"][0]
    assert result["ok"] is True
    assert result["markdown_filename"] == "Quarterly-Report.md"
    assert "Quarterly report" in result["markdown"]


def test_download_single_file_returns_markdown_attachment(client):
    response = client.post(
        "/api/convert/file",
        data=upload("notes.txt", b"plain text notes"),
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    assert response.mimetype == "text/markdown"
    assert "notes.md" in response.headers["Content-Disposition"]
    assert b"plain text notes" in response.data


def test_download_multiple_files_returns_zip(client):
    data = {
        "files": [
            (io.BytesIO(b"a,b\n1,2\n"), "one.csv"),
            (io.BytesIO(b"<h1>Two</h1>"), "two.html"),
        ]
    }
    response = client.post(
        "/api/convert/file", data=data, content_type="multipart/form-data"
    )
    assert response.status_code == 200
    assert response.mimetype == "application/zip"

    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        assert sorted(archive.namelist()) == ["one.md", "two.md"]


def test_zip_is_expanded_without_leaking_server_paths(client):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "inside the archive")
        archive.writestr("data.csv", "a,b\n1,2\n")

    payload = client.post(
        "/api/convert",
        data=upload("archive.zip", buffer.getvalue()),
        content_type="multipart/form-data",
    ).get_json()

    result = payload["results"][0]
    assert result["ok"] is True
    assert "inside the archive" in result["markdown"]
    assert "| a | b |" in result["markdown"]
    # The server's temporary directory must never appear in the output.
    assert "/tmp/" not in result["markdown"]
    assert "markitdown-" not in result["markdown"]


def test_duplicate_names_are_kept_distinct(client):
    data = {
        "files": [
            (io.BytesIO(b"<h1>One</h1>"), "report.html"),
            (io.BytesIO(b"first,second\n1,2\n"), "report.csv"),
        ]
    }
    payload = client.post(
        "/api/convert", data=data, content_type="multipart/form-data"
    ).get_json()

    names = [r["markdown_filename"] for r in payload["results"]]
    assert names == ["report.md", "report-2.md"]


def test_unsupported_extension_is_rejected_per_file(client):
    payload = client.post(
        "/api/convert",
        data=upload("payload.exe", b"MZ\x90\x00"),
        content_type="multipart/form-data",
    ).get_json()

    result = payload["results"][0]
    assert result["ok"] is False
    assert ".exe" in result["error"]
    assert payload["summary"]["failed"] == 1


def test_no_file_is_a_400(client):
    response = client.post("/api/convert", data={}, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "error" in response.get_json()


def test_too_many_files_is_rejected(client):
    client.application.config["MAX_FILES"] = 1
    data = {
        "files": [
            (io.BytesIO(b"<h1>One</h1>"), "one.html"),
            (io.BytesIO(b"<h1>Two</h1>"), "two.html"),
        ]
    }
    response = client.post("/api/convert", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "limit is 1" in response.get_json()["error"]


def test_oversized_upload_returns_413_json(client):
    client.application.config["MAX_CONTENT_LENGTH"] = 128
    response = client.post(
        "/api/convert",
        data=upload("big.txt", b"x" * 4096),
        content_type="multipart/form-data",
    )
    assert response.status_code == 413
    assert "too large" in response.get_json()["error"].lower()


def test_bundle_endpoint_zips_markdown(client):
    response = client.post(
        "/api/bundle",
        data=json.dumps(
            {
                "files": [
                    {"filename": "a.md", "markdown": "# A"},
                    {"filename": "a.md", "markdown": "# A again"},
                ]
            }
        ),
        content_type="application/json",
    )
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        assert sorted(archive.namelist()) == ["a-2.md", "a.md"]
        assert archive.read("a.md").decode() == "# A"


def test_bundle_rejects_empty_payload(client):
    response = client.post("/api/bundle", json={"files": []})
    assert response.status_code == 400


def test_api_404_is_json(client):
    response = client.get("/api/nope")
    assert response.status_code == 404
    assert response.get_json()["error"] == "Not found."


@pytest.mark.parametrize(
    ("given", "expected"),
    [
        ("report.pdf", "report.md"),
        ("My Report (final).docx", "My-Report-final.md"),
        ("../../etc/passwd", "passwd.md"),
        ("résumé.pdf", "resume.md"),
        ("...", "document.md"),
        ("", "document.md"),
    ],
)
def test_markdown_filename_is_safe(given, expected):
    assert markdown_filename_for(given) == expected


def test_safe_stem_truncates_long_names():
    assert len(safe_stem("x" * 500)) == 100
